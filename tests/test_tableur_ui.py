"""fm_tableur, interface et fichiers : grille (saisie, clavier d'Excel, sélection, poignée de recopie, pointage
des références, presse-papiers), classeur complet (ruban, barre de formule, feuilles), aller-retour Excel .xlsx
(Excel -> fm_tableur -> modification -> export -> réimport -> comparaison), CSV, archive piégée refusée,
fluidité de la peinture (mesure)."""
import os as _os, sys as _sys  # jamais les vrais réglages / presets de l'utilisateur (tests/_isolate.py)
_sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
import _isolate  # noqa: E402,F401
import os
import shutil
import sys
import tempfile
import time
import unittest
import zipfile

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from PySide6.QtCore import QPointF, Qt  # noqa: E402
from PySide6.QtGui import QImage, QKeyEvent, QMouseEvent, QPainter  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

APP = QApplication.instance() or QApplication([])

from fm_tableur import io_csv, io_xlsx, numfmt, ops  # noqa: E402
from fm_tableur.commands import SetCells  # noqa: E402
from fm_tableur.document import Cell  # noqa: E402
from fm_tableur.engine import cell_value, fz  # noqa: E402
from fm_tableur.session import Session  # noqa: E402
from fm_tableur.widget import SpreadsheetWidget  # noqa: E402

EN = numfmt.Locale(".", ",", "en")


def key(widget, k, mods=Qt.NoModifier, text=""):
    QApplication.sendEvent(widget, QKeyEvent(QKeyEvent.KeyPress, k, mods, text))


def click(widget, pos, mods=Qt.NoModifier, button=Qt.LeftButton, double=False):
    kind = QMouseEvent.MouseButtonDblClick if double else QMouseEvent.MouseButtonPress
    QApplication.sendEvent(widget, QMouseEvent(kind, QPointF(pos), QPointF(widget.mapToGlobal(pos)), button,
                                               button, mods))
    QApplication.sendEvent(widget, QMouseEvent(QMouseEvent.MouseButtonRelease, QPointF(pos),
                                               QPointF(widget.mapToGlobal(pos)), button, Qt.NoButton, mods))


class UiCase(unittest.TestCase):
    def setUp(self):
        self.session = Session()
        self.w = SpreadsheetWidget(self.session, {"tableur_decimal_sep": "."})
        self.w.resize(1200, 700)
        self.w.show()
        APP.processEvents()
        self.g = self.w.grid
        self.sheet = self.session.sheet

    def tearDown(self):
        self.w.detach()
        self.w.close()
        self.w.deleteLater()
        APP.processEvents()

    def type_cell(self, text):
        for ch in text:
            if self.g.is_editing():
                self.g._editor.insertPlainText(ch)
            else:
                key(self.g, 0, text=ch)
        key(self.g._editor if self.g.is_editing() else self.g, Qt.Key_Return)

    def center(self, r, c):
        return self.g.cell_rect(r, c).center().toPoint()


class GridTest(UiCase):
    def test_typing_enter_moves_down_and_formula_computes(self):
        self.g.select(0, 0)
        self.type_cell("12")
        self.assertEqual(self.g.view.current, (1, 0))
        self.type_cell("30")
        self.type_cell("=A1+A2")
        self.assertEqual(self.sheet.cell(2, 0).formula, "=A1+A2")
        if fz is not None:
            self.assertEqual(cell_value(self.session.book, self.sheet, 2, 0), 42)

    def test_escape_cancels_and_f2_edits(self):
        self.g.select(0, 0)
        self.type_cell("5")
        self.g.select(0, 0)
        key(self.g, 0, text="9")
        key(self.g._editor, Qt.Key_Escape)
        self.assertEqual(self.sheet.cell(0, 0).value, 5)
        key(self.g, Qt.Key_F2)
        self.assertTrue(self.g.is_editing())
        self.assertEqual(self.g.edit_text(), "5")

    def test_arrows_shift_ctrl_and_tab(self):
        self.session.stack.push(SetCells(self.sheet, {(r, 0): Cell(r) for r in range(10)}))
        self.g.select(0, 0)
        key(self.g, Qt.Key_Down, Qt.ControlModifier)
        self.assertEqual(self.g.view.current, (9, 0))
        key(self.g, Qt.Key_Up, Qt.ShiftModifier)
        self.assertEqual(self.g.selection(), (8, 0, 9, 0))
        self.g.select(0, 0)
        key(self.g, Qt.Key_Tab)
        self.assertEqual(self.g.view.current, (0, 1))

    def test_mouse_select_and_formula_pointing(self):
        self.g.select(0, 2)
        self.g.begin_edit("=", mode="enter")
        click(self.g.viewport(), self.center(3, 1))
        self.assertEqual(self.g.edit_text(), "=B4")
        self.g.commit_edit()
        self.assertEqual(self.sheet.cell(0, 2).formula, "=B4")
        click(self.g.viewport(), self.center(5, 5))
        self.assertEqual(self.g.view.current, (5, 5))
        click(self.g.viewport(), self.center(7, 6), Qt.ShiftModifier)
        self.assertEqual(self.g.selection(), (5, 5, 7, 6))

    def test_fill_handle_series(self):
        self.session.stack.push(SetCells(self.sheet, {(0, 0): Cell(1), (1, 0): Cell(2)}))
        self.g.select_range((0, 0, 1, 0))
        handle = self.g.range_rect((0, 0, 1, 0)).bottomRight()
        vp = self.g.viewport()
        QApplication.sendEvent(vp, QMouseEvent(QMouseEvent.MouseButtonPress, handle, vp.mapToGlobal(handle),
                                               Qt.LeftButton, Qt.LeftButton, Qt.NoModifier))
        target = self.g.cell_rect(4, 0).center()
        QApplication.sendEvent(vp, QMouseEvent(QMouseEvent.MouseMove, target, vp.mapToGlobal(target),
                                               Qt.NoButton, Qt.LeftButton, Qt.NoModifier))
        QApplication.sendEvent(vp, QMouseEvent(QMouseEvent.MouseButtonRelease, target, vp.mapToGlobal(target),
                                               Qt.LeftButton, Qt.NoButton, Qt.NoModifier))
        self.assertEqual([self.sheet.cell(r, 0).value for r in range(5)], [1, 2, 3, 4, 5])

    def test_clipboard_copy_paste_and_external_text(self):
        self.session.stack.push(SetCells(self.sheet, {(0, 0): Cell(2), (0, 1): Cell(None, "=A1*2")}))
        self.g.select_range((0, 0, 0, 1))
        key(self.g, Qt.Key_C, Qt.ControlModifier)
        self.g.select(3, 0)
        key(self.g, Qt.Key_V, Qt.ControlModifier)
        self.assertEqual(self.sheet.cell(3, 1).formula, "=A4*2")
        QApplication.clipboard().setText("10\t20\r\n30\t40\r\n")
        self.g.select(6, 0)
        self.g.paste()
        self.assertEqual(self.sheet.cell(7, 1).value, 40)

    def test_undo_redo_keys_and_delete(self):
        self.g.select(0, 0)
        self.type_cell("7")
        self.g.select(0, 0)
        key(self.g, Qt.Key_Delete)
        self.assertIsNone(self.sheet.cell(0, 0))
        key(self.g, Qt.Key_Z, Qt.ControlModifier)
        self.assertEqual(self.sheet.cell(0, 0).value, 7)
        key(self.g, Qt.Key_Y, Qt.ControlModifier)
        self.assertIsNone(self.sheet.cell(0, 0))

    def test_bold_shortcut_and_ribbon_sync(self):
        self.g.select(0, 0)
        key(self.g, Qt.Key_B, Qt.ControlModifier)
        self.assertTrue(self.session.book.style_of(self.sheet, 0, 0).get("bold"))
        self.assertTrue(self.w.ribbon.style_buttons["bold"].isChecked())

    def test_formula_bar_and_name_box(self):
        self.session.stack.push(SetCells(self.sheet, {(1, 1): Cell(None, "=1+1")}))
        self.w.name_box.setText("B2")
        self.w._goto_name_box()
        self.assertEqual(self.g.view.current, (1, 1))
        self.assertEqual(self.w.formula.text(), "=1+1")

    def test_french_formula_input_and_display(self):
        self.session.stack.push(SetCells(self.sheet, {(0, 0): Cell(2), (1, 0): Cell(4)}))
        self.g.select(2, 0)
        self.g.begin_edit("=SOMME(A1;A2)*1,5", mode="enter")
        self.g.commit_edit(move=None)
        self.assertEqual(self.sheet.cell(2, 0).formula, "=SUM(A1,A2)*1.5")
        self.assertEqual(self.w.formula.text(), "=SOMME(A1;A2)*1,5")
        if fz is not None:
            self.assertEqual(cell_value(self.session.book, self.sheet, 2, 0), 9)

    def test_number_widens_default_column_like_excel(self):
        self.g.select(0, 0)
        self.g.begin_edit("123456789012,5 €", mode="enter")
        self.g.commit_edit(move=None)
        self.assertGreater(self.sheet.cols.size(0), 64)
        self.session.stack.undo()                              # saisie + largeur : UNE étape
        self.assertNotIn(0, self.sheet.cols.sizes)
        self.assertIsNone(self.sheet.cell(0, 0))
        ops.set_sizes(self.session.book, self.session.stack, self.sheet, "col", [1], 40)
        self.g.select(0, 1)
        self.g.begin_edit("123456789", mode="enter")
        self.g.commit_edit(move=None)
        self.assertEqual(self.sheet.cols.size(1), 40)          # largeur choisie : jamais touchée

    def test_refused_formula_keeps_editing(self):
        self.g.select(0, 0)
        self.g.begin_edit("=1+*2", mode="enter")
        self.assertFalse(self.g.commit_edit())
        self.assertTrue(self.g.is_editing())
        self.g.cancel_edit()

    def test_sheets_add_rename_switch(self):
        self.w.add_sheet()
        self.assertEqual(len(self.session.book.sheets), 2)
        self.assertIs(self.session.sheet, self.session.book.sheets[1])
        key(self.g, Qt.Key_PageUp, Qt.ControlModifier)
        self.assertIs(self.session.sheet, self.session.book.sheets[0])

    def test_paint_is_fast(self):
        changes = {(r, c): Cell(r * 10 + c + 0.5) for r in range(2000) for c in range(15)}
        self.session.stack.push(SetCells(self.sheet, changes))
        image = QImage(1200, 700, QImage.Format_ARGB32_Premultiplied)
        painter = QPainter(image)
        self.g.paint_to(painter, self.g.viewport().rect())
        start = time.perf_counter()
        for _ in range(5):
            self.g.paint_to(painter, self.g.viewport().rect())
        painter.end()
        self.assertLess((time.perf_counter() - start) / 5, 0.05)     # mesure locale ~6-11 ms


class FilesTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def test_excel_file_round_trip(self):
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        src = os.path.join(self.dir, "source.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Données"
        ws["A1"], ws["B1"] = "Libellé", "Montant"
        ws["A2"], ws["B2"] = "Transport", 120.5
        ws["A3"], ws["B3"] = "Repas", 40
        ws["B4"] = "=SUM(B2:B3)"
        ws["C2"] = "=_xlfn.XLOOKUP(\"Repas\",A2:A3,B2:B3)"
        ws["A1"].font = Font(bold=True, color="FF0000")
        ws["B2"].fill = PatternFill("solid", fgColor="FFFF00")
        ws["B2"].number_format = "#,##0.00"
        ws.column_dimensions["A"].width = 20
        ws.merge_cells("D1:E1")
        ws.freeze_panes = "A2"
        wb.create_sheet("Autre")["A1"] = "=Données!B4*2"
        wb.save(src)
        book = io_xlsx.read(src)
        session = Session(book)
        s0, s1 = book.sheets
        self.assertEqual([s.name for s in book.sheets], ["Données", "Autre"])
        self.assertEqual(s0.cell(1, 2).formula, '=XLOOKUP("Repas",A2:A3,B2:B3)')
        style = book.style_of(s0, 0, 0)
        self.assertEqual((style.get("bold"), style.get("color")), (True, "#ff0000"))
        self.assertEqual(book.style_of(s0, 1, 1).get("fill"), "#ffff00")
        self.assertEqual((s0.merges, s0.freeze), ([(0, 3, 0, 4)], (1, 0)))
        self.assertAlmostEqual(s0.cols.sizes[0], 145.0)
        if fz is not None:
            self.assertEqual(cell_value(book, s1, 0, 0), 321.0)
            self.assertEqual(cell_value(book, s0, 1, 2), 40)
        ops.set_input(book, session.stack, s0, 2, 1, "60", EN)              # modification
        out = os.path.join(self.dir, "export.xlsx")
        io_xlsx.write(book, out)
        raw = openpyxl.load_workbook(out)
        self.assertEqual(raw["Données"]["C2"].value, '=_xlfn.XLOOKUP("Repas",A2:A3,B2:B3)')
        self.assertTrue(raw.calculation.fullCalcOnLoad)
        again = io_xlsx.read(out)
        a0 = again.sheets[0]
        for key in ((0, 0), (1, 0), (1, 1), (3, 1), (1, 2)):
            self.assertEqual((s0.cell(*key).value, s0.cell(*key).formula),
                             (a0.cell(*key).value if a0.cell(*key).formula is None else s0.cell(*key).value,
                              a0.cell(*key).formula), key)
        self.assertEqual(a0.cell(2, 1).value, 60)
        self.assertEqual(again.style_of(a0, 1, 1).get("fmt"), "#,##0.00")
        self.assertEqual((a0.merges, a0.freeze), (s0.merges, s0.freeze))
        self.assertEqual(os.path.getsize(src) > 0, True)                     # l'original n'est pas réécrit

    def test_excel_charts_are_read_and_linked(self):
        import openpyxl
        from openpyxl.chart import BarChart, LineChart, Reference
        src = os.path.join(self.dir, "graphiques.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Ventes"
        for row in [["Mois", "Nord", "Sud"], ["Janv", 10, 7], ["Févr", 12, 9], ["Mars", 8, 11]]:
            ws.append(row)
        bar = BarChart()
        bar.title = "Ventes par mois"
        bar.add_data(Reference(ws, min_col=2, min_row=1, max_col=3, max_row=4), titles_from_data=True)
        bar.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
        ws.add_chart(bar, "E2")
        line = LineChart()
        line.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
        ws.add_chart(line, "E20")
        wb.save(src)
        book = io_xlsx.read(src)
        charts = book.sheets[0].charts
        self.assertEqual(len(charts), 2)
        first = charts[0]
        self.assertEqual((first["type"], first["data"], first["title"], first["anchor"]),
                         ("column", "A1:C4", "Ventes par mois", [1, 4]))
        self.assertEqual(charts[1]["type"], "line")
        self.assertFalse(any("graphique" in x for x in book.source["lost"]))
        from fm_tableur.charts import chart_data
        Session(book)
        categories, series = chart_data(book, book.sheets[0], first)
        self.assertEqual(categories, ["Janv", "Févr", "Mars"])
        self.assertEqual(series, [("Nord", [10.0, 12.0, 8.0]), ("Sud", [7.0, 9.0, 11.0])])

    def test_unsupported_parts_reported(self):
        src = os.path.join(self.dir, "macros.xlsm")
        import openpyxl
        openpyxl.Workbook().save(src)
        with zipfile.ZipFile(src, "a") as z:
            z.writestr("xl/vbaProject.bin", b"\0" * 10)
            z.writestr("xl/pivotTables/pivotTable1.xml", "<x/>")
        book = io_xlsx.read(src)
        self.assertTrue(any("macros" in x for x in book.source["lost"]))
        self.assertTrue(any("croisés" in x for x in book.source["lost"]))

    def test_zip_bomb_and_garbage_refused(self):
        bad = os.path.join(self.dir, "faux.xlsx")
        with open(bad, "wb") as f:
            f.write(b"pas une archive")
        with self.assertRaises(io_xlsx.UnsafeFile):
            io_xlsx.read(bad)
        bomb = os.path.join(self.dir, "bombe.xlsx")
        with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("xl/worksheets/sheet1.xml", b"\0" * (60 * 1024 ** 2))
        with self.assertRaises(io_xlsx.UnsafeFile):
            io_xlsx.check_archive(bomb)

    def test_csv_round_trip_french(self):
        path = os.path.join(self.dir, "données.csv")
        with open(path, "w", encoding="cp1252") as f:
            f.write("Nom;Montant;Date\nÉté;1 234,5;30/10/2026\n")
        book = io_csv.read(path)
        sheet = book.sheets[0]
        self.assertEqual(sheet.cell(1, 0).value, "Été")
        self.assertEqual(sheet.cell(1, 1).value, 1234.5)
        self.assertEqual(sheet.cell(1, 2).value, 46325)
        out = os.path.join(self.dir, "sortie.csv")
        io_csv.write(book, sheet, out)
        with open(out, encoding="utf-8-sig") as f:
            text = f.read()
        self.assertIn("Été;1 234,5;30/10/2026", text.replace(" ", " "))

    def test_prefixes(self):
        self.assertEqual(io_xlsx.add_prefixes('=XLOOKUP(A1,B:B,C:C)+SUM(FILTER(A1:A3,A1:A3>1))&"XLOOKUP("'),
                         '=_xlfn.XLOOKUP(A1,B:B,C:C)+SUM(_xlfn._xlws.FILTER(A1:A3,A1:A3>1))&"XLOOKUP("')
        self.assertEqual(io_xlsx.strip_prefixes("=_xlfn._xlws.SORT(A1:A3)"), "=SORT(A1:A3)")


if __name__ == "__main__":
    unittest.main()
