"""fm_tableur dans le Board : sélecteur de modules, élément « sheet » (création, aperçu, Ctrl + glisser), édition sur
place (souris et clavier laissés au classeur, Échap pour sortir), enregistrement automatique et réouverture du
.mro, mode Standard sur la MÊME session (document, annulation, feuille, sélection gardés, rien de rechargé),
compatibilité (ancien fichier, champs inconnus, autres modules intacts)."""
import os as _os, sys as _sys  # jamais les vrais réglages / presets de l'utilisateur (tests/_isolate.py)
_sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
import _isolate  # noqa: E402,F401
import json
import os
import shutil
import sys
import tempfile
import time
import unittest
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from PySide6.QtCore import QEvent, QPointF, Qt  # noqa: E402
from PySide6.QtGui import QKeyEvent, QMouseEvent  # noqa: E402
from PySide6.QtWidgets import QApplication, QWidget  # noqa: E402

APP = QApplication.instance() or QApplication([])

from fm_board import serialization  # noqa: E402
from fm_board import ui as board_ui  # noqa: E402
from fm_tableur import session as S  # noqa: E402
from fm_tableur.board_item import SheetItem  # noqa: E402
from fm_tableur.engine import cell_value, fz  # noqa: E402

_PANES = []


def settle(seconds=0.05):
    end = time.monotonic() + seconds
    while time.monotonic() < end:
        APP.processEvents()
        time.sleep(0.003)


def key(widget, k, mods=Qt.NoModifier, text=""):
    QApplication.sendEvent(widget, QKeyEvent(QKeyEvent.KeyPress, k, mods, text))
    QApplication.sendEvent(widget, QKeyEvent(QKeyEvent.KeyRelease, k, mods, text))


class BoardCase(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="tableur_test_"))
        self.path = serialization.create_board_file(self.tmp / "essai.mro")
        self.host = QWidget()
        self.host.resize(1400, 900)
        self.pane = board_ui.BoardPane(self.host)
        _PANES.append(self.pane)
        self.assertTrue(self.pane.open(self.path))
        self.pane.resize(1400, 900)
        self.host.show()
        self.pane.show()
        self.pane.set_module("tableur")
        self.ctl = self.pane.tableur
        self.scene = self.pane.scene
        self.view = self.pane.view
        settle()

    def tearDown(self):
        from PySide6.QtCore import QCoreApplication
        for pane in _PANES:
            try:
                std = getattr(pane, "_tableur_standard", None)
                if std is not None:
                    std.close_pane()
                for item in pane.scene.board_items():
                    if isinstance(item, SheetItem):
                        item.stop_edit()
                pane._autosave.stop()
                pane.deleteLater()
            except RuntimeError:
                pass
        _PANES.clear()
        self.host.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
        shutil.rmtree(self.tmp, ignore_errors=True)

    def view_point(self, item, x, y):
        return self.view.mapFromScene(item.mapToScene(QPointF(x, y)))


class ModuleTest(BoardCase):
    def test_selector_and_new_sheet(self):
        self.assertIn("tableur", self.pane._module_buttons)
        self.assertTrue(self.ctl.active and self.ctl.bar.isVisible())
        item = self.ctl.new_sheet()
        self.assertIsInstance(item, SheetItem)
        self.assertEqual(item.MODULE, "tableur")
        self.scene.undo.undo()
        self.assertIsNone(self.scene.item(item.item_id))
        self.pane.set_module("miro")
        self.assertFalse(self.ctl.bar.isVisible())

    def test_fade_and_hide_like_other_modules(self):
        item = self.ctl.new_sheet()
        self.pane.toggle_module_look("tableur", "hide")
        self.assertFalse(item.isVisible())
        self.pane.toggle_module_look("tableur", "hide")
        self.assertTrue(item.isVisible())

    def test_move_needs_ctrl(self):
        item = self.ctl.new_sheet()
        settle()
        start = item.pos()
        p = self.view_point(item, 100, 60)
        vp = self.view.viewport()
        for kind, pos, buttons in ((QMouseEvent.MouseButtonPress, p, Qt.LeftButton),
                                   (QMouseEvent.MouseMove, p + QPointF(80, 40).toPoint(), Qt.LeftButton),
                                   (QMouseEvent.MouseButtonRelease, p + QPointF(80, 40).toPoint(), Qt.NoButton)):
            QApplication.sendEvent(vp, QMouseEvent(kind, QPointF(pos), QPointF(vp.mapToGlobal(pos)),
                                                   Qt.LeftButton, buttons, Qt.NoModifier))
        settle()
        self.assertEqual(item.pos(), start)                    # sans Ctrl : pas de déplacement


class EditTest(BoardCase):
    def test_edit_in_place_keyboard_and_autosave(self):
        item = self.ctl.new_sheet()
        item.start_edit()
        settle()
        self.assertTrue(item.is_editing())
        self.assertIs(self.scene.embedded_item, item)
        grid = item._widget.grid
        grid.select(0, 0)
        for ch in "21":
            key(grid, 0, text=ch) if not grid.is_editing() else grid._editor.insertPlainText(ch)
        key(grid._editor, Qt.Key_Return)
        grid.select(0, 1)
        grid.begin_edit("=A1*2", mode="enter")
        key(grid._editor, Qt.Key_Return)
        sheet = item.session.sheet
        self.assertEqual(sheet.cell(0, 0).value, 21)
        if fz is not None:
            self.assertEqual(cell_value(item.session.book, sheet, 0, 1), 42)
        self.assertTrue(self.pane.is_dirty())
        key(grid, Qt.Key_Escape)                               # rien à annuler : sortie de l'édition
        settle()
        self.assertFalse(item.is_editing())
        self.assertTrue(self.pane.save())
        data = json.loads(self.path.read_text(encoding="utf-8"))
        saved = [i for i in data["items"] if i["type"] == "sheet"][0]
        cells = saved["workbook"]["sheets"][0]["cells"]
        self.assertIn([0, 0, 21], cells)
        self.assertIn([0, 1, 42, "=A1*2"], [c[:4] for c in cells])
        reopened = board_ui.BoardPane(self.host)
        _PANES.append(reopened)
        self.assertTrue(reopened.open(self.path))
        other = [i for i in reopened.scene.board_items() if isinstance(i, SheetItem)][0]
        self.assertIsNot(other.session, item.session)          # nouvelle ouverture = relu du fichier
        self.assertEqual(other.session.sheet.cell(0, 1).formula, "=A1*2")

    def test_board_keys_do_not_reach_board_while_editing(self):
        item = self.ctl.new_sheet()
        item.start_edit()
        settle()
        grid = item._widget.grid
        grid.select(2, 2)
        key(grid, Qt.Key_Delete)                               # Suppr : efface la cellule, pas l'élément
        self.assertIsNotNone(self.scene.item(item.item_id))

    def test_undo_of_board_move_keeps_workbook(self):
        item = self.ctl.new_sheet()
        self.scene.change_items([item], lambda i: i.data.update(x=i.data["x"] + 50), "Déplacer")
        from fm_tableur import ops, numfmt
        ops.set_input(item.session.book, item.session.stack, item.session.sheet, 0, 0, "5", numfmt.DEFAULT_LOCALE)
        self.scene.undo.undo()                                 # annule le déplacement du Board
        self.assertEqual(item.session.sheet.cell(0, 0).value, 5)


class StandardTest(BoardCase):
    def test_same_session_both_ways(self):
        from fm_tableur import numfmt, ops
        item = self.ctl.new_sheet()
        session = item.session
        ops.add_sheet(session.book, session.stack)
        second = session.book.sheets[1]
        session.set_active(second)
        ops.set_input(session.book, session.stack, second, 3, 2, "abc", numfmt.DEFAULT_LOCALE)
        session.view(second).current = (3, 2)
        depth = len(session.stack)
        item.open_standard()
        settle()
        std = self.pane._tableur_standard
        self.assertIsNotNone(std)
        self.assertTrue(std.isVisible())
        self.assertIs(std.widget.session, session)             # même document, rien de rechargé
        self.assertIs(std.widget.grid.sheet, second)
        self.assertEqual(std.widget.grid.view.current, (3, 2))
        self.assertEqual(std.widget.formula.text(), "abc")
        std.widget.grid.select(4, 2)
        std.widget.grid.begin_edit("xyz", mode="enter")
        std.widget.grid.commit_edit()
        self.assertEqual(len(session.stack), depth + 1)
        std.close_pane()
        settle()
        self.assertIsNone(self.pane._tableur_standard)
        self.assertIs(item.session, session)
        self.assertEqual(session.view(second).current, (5, 2))
        session.stack.undo()                                   # l'annulation suit le document
        self.assertIsNone(second.cell(4, 2))


class CompatibilityTest(BoardCase):
    def test_old_board_and_unknown_fields_kept(self):
        old = {"format": "pipeline-board", "version": 2, "page": {}, "view": {"zoom": 1.0, "center": [0, 0]},
               "items": [{"id": "s1", "type": "sticky", "x": 1, "y": 2, "w": 200, "h": 200, "text": "hé"},
                         {"id": "t1", "type": "sheet", "x": 0, "y": 300, "w": 500, "h": 300, "futur": 7,
                          "workbook": {"format": "fm_tableur", "version": 1, "sheets": [
                              {"name": "A", "cells": [[0, 0, 1], [0, 1, None, "=A1+1"]], "inconnu": 3}],
                              "cle_future": {"x": 1}}}]}
        path = self.tmp / "ancien.mro"
        path.write_text(json.dumps(old), encoding="utf-8")
        other = board_ui.BoardPane(self.host)
        _PANES.append(other)
        self.assertTrue(other.open(path))
        item = other.scene.item("t1")
        if fz is not None:
            self.assertEqual(cell_value(item.session.book, item.session.sheet, 0, 1), 2)
        other._set_dirty(True)
        self.assertTrue(other.save())
        again = json.loads(path.read_text(encoding="utf-8"))
        sheet = [i for i in again["items"] if i["id"] == "t1"][0]
        self.assertEqual(sheet["futur"], 7)
        self.assertEqual(sheet["workbook"]["cle_future"], {"x": 1})
        self.assertEqual(sheet["workbook"]["sheets"][0]["inconnu"], 3)
        self.assertEqual([i["text"] for i in again["items"] if i["type"] == "sticky"], ["hé"])

    def test_sessions_forgotten_with_scene(self):
        item = self.ctl.new_sheet()
        key = item.session.key
        self.assertIs(S.get(key), item.session)


if __name__ == "__main__":
    unittest.main()


class ExplorerFileTest(BoardCase):
    """Classeur ouvert depuis l'explorateur : board en mémoire, mode Standard, original jamais réécrit sans
    demande ; « Enregistrer » confirmé = copie .bak puis remplacement."""

    def test_open_edit_save_with_backup(self):
        from unittest import mock
        import openpyxl
        from PySide6.QtWidgets import QMessageBox
        from fm_tableur import files_mode, io_xlsx, numfmt, ops
        src = self.tmp / "budget.xlsx"
        wb = openpyxl.Workbook()
        wb.active["A1"] = 10
        wb.active["A2"] = "=A1*3"
        wb.save(src)
        before = src.read_bytes()
        other = board_ui.BoardPane(self.host)
        _PANES.append(other)
        self.assertTrue(other.open_import(files_mode.read_as_board(src), src))
        item = files_mode.attach(other, src)
        settle()
        std = other._tableur_standard
        self.assertIsNotNone(std)
        self.assertIs(std.widget.session, item.session)
        session = item.session
        ops.set_input(session.book, session.stack, session.sheet, 0, 0, "20", numfmt.DEFAULT_LOCALE)
        self.assertTrue(other.is_dirty())
        self.assertEqual(src.read_bytes(), before)               # rien d'écrit sans le demander
        with mock.patch.object(QMessageBox, "exec", lambda box: box.setProperty("x", 1)), \
                mock.patch.object(QMessageBox, "clickedButton", lambda box: box.buttons()[0]):
            self.assertTrue(files_mode.save(other, item, src))  # « Remplacer »
        self.assertEqual((self.tmp / "budget.xlsx.bak").read_bytes(), before)
        again = io_xlsx.read(src)
        self.assertEqual(again.sheets[0].cell(0, 0).value, 20)
        self.assertEqual(again.sheets[0].cell(1, 0).formula, "=A1*3")
        self.assertFalse(other.is_dirty())
        std.close_pane()


class SettingsTest(BoardCase):
    def test_rows_declared_and_applied_live(self):
        import settings_store
        from fm_tableur.settings import TABLEUR_TREE
        from settings_catalog import CATALOG
        keys = {spec.key for node in TABLEUR_TREE for spec in CATALOG[node.group]}
        self.assertIn("tableur_formula_lang", keys)
        self.assertIn("tableur_font_family", keys)
        self.assertTrue(keys <= set(settings_store.DEFAULT_SETTINGS))       # clés déclarées (sinon ignorées)
        for key in keys:                                                    # défauts identiques des deux côtés
            spec = next(s for node in TABLEUR_TREE for s in CATALOG[node.group] if s.key == key)
            self.assertEqual(spec.default, settings_store.DEFAULT_SETTINGS[key], key)
        from fm_tableur import document
        self.pane.apply_settings({**self.pane._settings, "tableur_font_family": "Arial", "tableur_font_size": 14,
                                  "tableur_calc_mode": "manual"})
        try:
            fresh = self.ctl.new_sheet().session.book                   # nouveau classeur : réglages
            self.assertEqual((fresh.default_font, fresh.calc_mode), (("Arial", 14.0), "manual"))
        finally:
            self.pane.apply_settings({**self.pane._settings, "tableur_font_family": "", "tableur_font_size": 11,
                                      "tableur_calc_mode": "auto"})
        self.assertEqual((document.new_workbook().default_font, document.new_workbook().calc_mode),
                         (("Calibri", 11.0), "auto"))
        item = self.ctl.new_sheet()
        session = item.session
        from fm_tableur.commands import SetCells
        from fm_tableur.document import Cell
        session.stack.push(SetCells(session.sheet, {(0, 0): Cell(None, "=SUM(1,2)")}))
        item.open_standard()
        settle()
        grid = self.pane._tableur_standard.widget.grid
        self.assertEqual(grid.current_input_text(), "=SOMME(1;2)")
        self.pane.apply_settings({**self.pane._settings, "tableur_formula_lang": "en"})
        self.assertEqual(grid.current_input_text(), "=SUM(1,2)")
