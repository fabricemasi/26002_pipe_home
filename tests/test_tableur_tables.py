"""fm_tableur, tableaux structurés : création (entêtes détectées / ajoutées), références structurées traduites pour le
moteur (T[Col], T[[Col1]:[Col2]], #All / #Data / #Headers / #Totals, [@Col], T[@Col], nom seul), recalcul quand le
tableau change, extension automatique, colonne calculée, ligne des totaux, renommer, convertir en plage, insertion de
lignes, annulation, format natif, aller-retour Excel (tableaux et forme longue [#This Row])."""
import os as _os, sys as _sys  # jamais les vrais réglages / presets de l'utilisateur (tests/_isolate.py)
_sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
import _isolate  # noqa: E402,F401
import os
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fm_tableur import io_native, numfmt, ops, structured as T, tables_ops  # noqa: E402
from fm_tableur import values as V  # noqa: E402
from fm_tableur.document import new_workbook  # noqa: E402
from fm_tableur.engine import cell_value, fz  # noqa: E402
from fm_tableur.session import Session  # noqa: E402

EN = numfmt.Locale(".", ",", "en")


def make():
    s = Session(new_workbook(1))
    sheet = s.book.sheets[0]
    for r, row in enumerate([["Poste", "Montant", "Qte"], ["Transport", 10, 2], ["Repas", 20, 3],
                             ["Hôtel", 30, 4]]):
        for c, v in enumerate(row):
            ops.set_input(s.book, s.stack, sheet, r, c, str(v), EN)
    return s, sheet


def put(s, sheet, r, c, text):
    ops.set_input(s.book, s.stack, sheet, r, c, text, EN)


def val(s, sheet, r, c):
    return cell_value(s.book, sheet, r, c)


class ResolveTest(unittest.TestCase):
    def setUp(self):
        self.s, self.sheet = make()
        self.table = tables_ops.create(self.s.book, self.s.stack, self.sheet, (0, 0, 3, 2))

    def r(self, formula, row=10, col=6):
        return T.resolve(formula, self.s.book, self.sheet, row, col)

    def test_creation_detects_header(self):
        self.assertEqual((self.table["name"], self.table["ref"], self.table["header"]), ("Tableau1", (0, 0, 3, 2), True))
        self.assertEqual(T.columns(self.s.book, self.sheet, self.table), ["Poste", "Montant", "Qte"])

    def test_translation(self):
        n = "Feuil1!"
        self.assertEqual(self.r("=SUM(Tableau1[Montant])"), f"=SUM({n}$B$2:$B$4)")
        self.assertEqual(self.r("=SUM(Tableau1[[Montant]:[Qte]])"), f"=SUM({n}$B$2:$C$4)")
        self.assertEqual(self.r("=ROWS(Tableau1[#All])"), f"=ROWS({n}$A$1:$C$4)")
        self.assertEqual(self.r("=COUNTA(Tableau1[#Headers])"), f"=COUNTA({n}$A$1:$C$1)")
        self.assertEqual(self.r("=SUM(Tableau1[[#Headers],[#Data],[Qte]])"), f"=SUM({n}$C$1:$C$4)")
        self.assertEqual(self.r("=ROWS(Tableau1)+ROWS(Tableau1[])"), f"=ROWS({n}$A$2:$C$4)+ROWS({n}$A$2:$C$4)")
        self.assertEqual(self.r("=[@Montant]*[@Qte]", 2, 0), f"={n}$B$3*{n}$C$3")
        self.assertEqual(self.r("=Tableau1[@Montant]", 2, 9), f"={n}$B$3")
        self.assertEqual(self.r("=Tableau1[[#This Row],[Qte]]", 3, 9), f"={n}$C$4")
        self.assertEqual(self.r('=Tableau1[Inconnue]&"Tableau1[Montant]"'), '=#REF!&"Tableau1[Montant]"')
        self.assertEqual(self.r("=Absent[Montant]"), "=#REF!")
        self.assertEqual(self.r("=[@Montant]", 20, 1), "=#REF!")                 # hors de tout tableau
        self.assertEqual(self.r("=Tableau1[@Montant]", 20, 1), "=#VALUE!")       # ligne hors du tableau

    def test_french_keywords(self):
        self.assertEqual(self.r("=COUNTA(Tableau1[#En-têtes])"), "=COUNTA(Feuil1!$A$1:$C$1)")
        self.assertEqual(self.r("=Tableau1[[#Cette ligne],[Qte]]", 3, 9), "=Feuil1!$C$4")

    def test_escaped_column_names(self):
        put(self.s, self.sheet, 0, 2, "Prix [HT] #1")
        self.assertEqual(self.r("=SUM(Tableau1[Prix '[HT'] '#1])"), "=SUM(Feuil1!$C$2:$C$4)")


@unittest.skipIf(fz is None, "formualizer absent")
class EngineTablesTest(unittest.TestCase):
    def setUp(self):
        self.s, self.sheet = make()
        tables_ops.create(self.s.book, self.s.stack, self.sheet, (0, 0, 3, 2))

    def test_formulas_compute_and_follow_changes(self):
        s, sh = self.s, self.sheet
        put(s, sh, 6, 0, "=SUM(Tableau1[Montant])")
        put(s, sh, 7, 0, '=XLOOKUP("Repas",Tableau1[Poste],Tableau1[Montant])')
        put(s, sh, 8, 0, "=ROWS(Tableau1)")
        self.assertEqual((val(s, sh, 6, 0), val(s, sh, 7, 0), val(s, sh, 8, 0)), (60, 20, 3))
        put(s, sh, 2, 1, "25")                                   # donnée modifiée
        self.assertEqual(val(s, sh, 6, 0), 65)
        put(s, sh, 4, 0, "Divers")                               # ligne sous le tableau : extension auto
        self.assertEqual(T.find(s.book, "Tableau1")[1]["ref"], (0, 0, 4, 2))
        put(s, sh, 4, 1, "5")
        self.assertEqual((val(s, sh, 6, 0), val(s, sh, 8, 0)), (70, 4))
        put(s, sh, 0, 1, "Prix")                                 # entête renommée : les formules suivent (Excel)
        self.assertEqual(sh.cell(6, 0).formula, "=SUM(Tableau1[Prix])")
        self.assertEqual(sh.cell(7, 0).formula, '=XLOOKUP("Repas",Tableau1[Poste],Tableau1[Prix])')
        self.assertEqual(val(s, sh, 6, 0), 70)
        s.stack.undo()                                           # une seule étape : entête et formules
        self.assertEqual(sh.cell(6, 0).formula, "=SUM(Tableau1[Montant])")
        self.assertEqual(val(s, sh, 6, 0), 70)
        put(s, sh, 6, 1, "=SUM(Tableau1[Absente])")
        self.assertEqual(val(s, sh, 6, 1), V.XLError("#REF!"))

    def test_calculated_column_and_this_row(self):
        s, sh = self.s, self.sheet
        tables_ops.after_input  # (accroché à la saisie)
        put(s, sh, 0, 3, "Total")                                # colonne à droite : extension
        self.assertEqual(T.find(s.book, "Tableau1")[1]["ref"], (0, 0, 3, 3))
        put(s, sh, 1, 3, "=[@Montant]*[@Qte]")                    # colonne calculée : recopiée
        self.assertEqual([sh.cell(r, 3).formula for r in (1, 2, 3)], ["=[@Montant]*[@Qte]"] * 3)
        self.assertEqual([val(s, sh, r, 3) for r in (1, 2, 3)], [20, 60, 120])
        put(s, sh, 4, 0, "Divers")                               # nouvelle ligne : la colonne calculée la suit
        self.assertEqual(sh.cell(4, 3).formula, "=[@Montant]*[@Qte]")
        s.stack.undo()
        self.assertIsNone(sh.cell(4, 3))

    def test_totals_rename_insert_and_convert(self):
        s, sh = self.s, self.sheet
        table = T.find(s.book, "Tableau1")[1]
        tables_ops.toggle_totals(s.book, s.stack, sh, table)
        self.assertEqual((sh.cell(4, 0).value, sh.cell(4, 2).formula), ("Total", "=SUBTOTAL(109,Tableau1[Qte])"))
        self.assertEqual(val(s, sh, 4, 2), 9)
        put(s, sh, 7, 0, "=SUM(Tableau1[Montant])+Tableau1[[#Totals],[Qte]]")
        self.assertEqual(val(s, sh, 7, 0), 69)
        self.assertIsNone(tables_ops.rename(s.book, s.stack, sh, T.find(s.book, "Tableau1")[1], "Frais"))
        self.assertEqual(sh.cell(7, 0).formula, "=SUM(Frais[Montant])+Frais[[#Totals],[Qte]]")
        self.assertEqual(val(s, sh, 7, 0), 69)
        self.assertIsNotNone(tables_ops.rename(s.book, s.stack, sh, T.find(s.book, "Frais")[1], "A1"))
        ops.insert_lines(s.book, s.stack, sh, "row", 2, 1)       # ligne insérée DANS le tableau
        self.assertEqual(T.find(s.book, "Frais")[1]["ref"], (0, 0, 5, 2))
        put(s, sh, 2, 1, "100")
        self.assertEqual(val(s, sh, 8, 0), 169)
        tables_ops.convert_to_range(s.book, s.stack, sh, T.find(s.book, "Frais")[1])
        self.assertEqual(sh.tables, [])
        self.assertEqual(sh.cell(8, 0).formula, "=SUM($B$2:$B$5)+$C$6")
        self.assertEqual(val(s, sh, 8, 0), 169)
        s.stack.undo()
        self.assertEqual(sh.cell(8, 0).formula, "=SUM(Frais[Montant])+Frais[[#Totals],[Qte]]")
        self.assertEqual(val(s, sh, 8, 0), 169)

    def test_native_round_trip(self):
        put(self.s, self.sheet, 6, 0, "=SUM(Tableau1[Montant])")
        book, _a = io_native.from_dict(io_native.to_dict(self.s.book))
        other = Session(book)
        self.assertEqual(book.sheets[0].tables[0]["ref"], (0, 0, 3, 2))
        self.assertEqual(cell_value(other.book, book.sheets[0], 6, 0), 60)


class ExcelTablesTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def test_import_and_export(self):
        import openpyxl
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from fm_tableur import io_xlsx
        src = os.path.join(self.dir, "tableaux.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        for row in [["Poste", "Montant", "Qte", "Total"], ["A", 10, 2, None], ["B", 20, 3, None]]:
            ws.append(row)
        ws["D2"] = "=Ventes[[#This Row],[Montant]]*Ventes[[#This Row],[Qte]]"
        ws["D3"] = "=Ventes[[#This Row],[Montant]]*Ventes[[#This Row],[Qte]]"
        ws["F1"] = "=SUM(Ventes[Total])"
        table = Table(displayName="Ventes", ref="A1:D3")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium7", showRowStripes=True)
        ws.add_table(table)
        wb.save(src)
        book = io_xlsx.read(src)
        session = Session(book)
        sheet = book.sheets[0]
        self.assertEqual(sheet.tables[0]["name"], "Ventes")
        self.assertEqual((sheet.tables[0]["ref"], sheet.tables[0]["style"]), ((0, 0, 2, 3), "vert"))
        self.assertFalse(any("structurés" in x for x in book.source["lost"]))
        if fz is not None:
            self.assertEqual(cell_value(book, sheet, 0, 5), 80)
        ops.set_input(book, session.stack, sheet, 1, 3, "=[@Montant]+1", EN)   # forme courte d'Excel
        out = os.path.join(self.dir, "export.xlsx")
        io_xlsx.write(book, out)
        raw = openpyxl.load_workbook(out)
        ws2 = raw.active
        self.assertIn("Ventes", ws2.tables)
        self.assertEqual(ws2.tables["Ventes"].ref, "A1:D3")
        self.assertEqual(ws2["D2"].value, "=Ventes[[#This Row],[Montant]]+1")     # forme longue du fichier
        again = io_xlsx.read(out)
        self.assertEqual(again.sheets[0].tables[0]["ref"], (0, 0, 2, 3))


if __name__ == "__main__":
    unittest.main()


class CtrlTTest(unittest.TestCase):
    def test_ctrl_t_creates_table_from_current_region(self):
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtCore import Qt
        from PySide6.QtGui import QKeyEvent
        from PySide6.QtWidgets import QApplication
        app = QApplication.instance() or QApplication([])
        from fm_tableur.widget import SpreadsheetWidget
        s, sheet = make()
        w = SpreadsheetWidget(s)
        w.resize(900, 500)
        w.show()
        app.processEvents()
        try:
            w.grid.select(2, 1)
            QApplication.sendEvent(w.grid, QKeyEvent(QKeyEvent.KeyPress, Qt.Key_T, Qt.ControlModifier))
            self.assertEqual(sheet.tables[0]["ref"], (0, 0, 3, 2))
            self.assertEqual(w.grid.selection(), (0, 0, 3, 2))
            from PySide6.QtGui import QImage, QPainter
            image = QImage(900, 500, QImage.Format_ARGB32_Premultiplied)
            painter = QPainter(image)
            w.grid.paint_to(painter, w.grid.viewport().rect())         # style à bandes peint sans erreur
            painter.end()
            s.stack.undo()
            self.assertEqual(sheet.tables, [])
        finally:
            w.detach()
            w.deleteLater()


class TableFiltersTest(unittest.TestCase):
    def setUp(self):
        from fm_tableur import filters
        self.F = filters
        self.s, self.sheet = make()
        put(self.s, self.sheet, 4, 0, "Repas")
        put(self.s, self.sheet, 4, 1, "5")
        put(self.s, self.sheet, 4, 2, "1")
        self.table = tables_ops.create(self.s.book, self.s.stack, self.sheet, (0, 0, 4, 2))
        self.scope = self.F.scope_at(self.sheet, 0, 0)

    def test_buttons_values_criteria_and_clear(self):
        s, sh, F = self.s, self.sheet, self.F
        self.assertTrue(self.table["filter"])
        self.assertEqual((self.scope.kind, self.scope.first, self.scope.last), ("table", 1, 4))
        self.assertIsNone(F.scope_at(sh, 1, 0))                       # seulement la ligne d'entête
        self.assertEqual(F.values(s.book, sh, self.scope, 0), ["Hôtel", "Repas", "Transport"])
        F.apply(s.book, s.stack, sh, self.scope, 0, {"Repas"})
        self.assertEqual(sh.rows.hidden, {1, 3})
        F.apply(s.book, s.stack, sh, F.scope_at(sh, 0, 1), 1, {"20"})   # deux critères : ET
        self.assertEqual(sh.rows.hidden, {1, 3, 4})
        self.assertEqual(sh.tables[0]["filters"], {"0": {"values": ["Repas"]}, "1": {"values": ["20"]}})
        s.stack.undo()
        self.assertEqual(sh.rows.hidden, {1, 3})
        F.clear(s.book, s.stack, sh, F.scope_at(sh, 0, 0))
        self.assertEqual((sh.rows.hidden, sh.tables[0]["filters"]), (set(), {}))
        F.apply(s.book, s.stack, sh, F.scope_at(sh, 0, 0), 0, {"Hôtel"})
        F.set_table_buttons(s.book, s.stack, sh, sh.tables[0], False)  # boutons retirés : tout réaffiché
        self.assertEqual((sh.rows.hidden, F.scope_at(sh, 0, 0)), (set(), None))

    def test_criteria_follow_the_table_and_sheet_filter_still_works(self):
        s, sh, F = self.s, self.sheet, self.F
        F.apply(s.book, s.stack, sh, self.scope, 2, {"4"})
        ops.insert_lines(s.book, s.stack, sh, "col", 0, 2)              # tableau déplacé de deux colonnes
        scope = F.scope_at(sh, 0, 2)
        self.assertEqual(scope.criteria(sh), {4: {"values": ["4"]}})
        self.assertEqual(F.values(s.book, sh, scope, 4), ["1", "2", "3", "4"])
        other, osheet = make()
        ops.set_autofilter(other.book, other.stack, osheet, (0, 0, 3, 2))
        self.assertEqual(ops.filter_values(other.book, osheet, 0), ["Hôtel", "Repas", "Transport"])
        ops.apply_filter(other.book, other.stack, osheet, 0, {"Repas"})
        self.assertEqual(osheet.rows.hidden, {1, 3})

    def test_excel_round_trip_of_filters(self):
        import openpyxl
        from fm_tableur import io_xlsx
        s, sh, F = self.s, self.sheet, self.F
        F.apply(s.book, s.stack, sh, self.scope, 0, {"Repas", ""})
        tmp = tempfile.mkdtemp()
        try:
            out = os.path.join(tmp, "filtres.xlsx")
            io_xlsx.write(s.book, out)
            raw = openpyxl.load_workbook(out).active.tables["Tableau1"]
            self.assertEqual(raw.autoFilter.ref, "A1:C5")
            again = io_xlsx.read(out)
            table = again.sheets[0].tables[0]
            self.assertTrue(table["filter"])
            self.assertEqual(table["filters"], {"0": {"values": ["", "Repas"]}})
            self.assertEqual(again.sheets[0].rows.hidden, {1, 3})
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_grid_button_opens_the_table_menu(self):
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtCore import QPointF, Qt
        from PySide6.QtGui import QMouseEvent
        from PySide6.QtWidgets import QApplication
        app = QApplication.instance() or QApplication([])
        from fm_tableur.widget import SpreadsheetWidget
        w = SpreadsheetWidget(self.s)
        w.resize(900, 500)
        w.show()
        app.processEvents()
        got = []
        w.grid.filterMenuWanted.disconnect()
        w.grid.filterMenuWanted.connect(lambda r, c, _p: got.append((r, c)))
        try:
            self.assertEqual([(r, c) for r, c, _on in w.grid.filter_buttons()], [(0, 0), (0, 1), (0, 2)])
            center = w.grid.filter_button_rect(0, 1).center()
            vp = w.grid.viewport()
            QApplication.sendEvent(vp, QMouseEvent(QMouseEvent.MouseButtonPress, center, vp.mapToGlobal(center),
                                                   Qt.LeftButton, Qt.LeftButton, Qt.NoModifier))
            self.assertEqual(got, [(0, 1)])
            self.assertNotEqual(w.grid.view.current, (0, 1))           # le clic n'a pas choisi la cellule
            w._filter_menu(0, 0, vp.mapToGlobal(QPointF(10, 10).toPoint()))   # menu construit sans erreur
        finally:
            w.detach()
            w.deleteLater()


class StylesAndHandleTest(unittest.TestCase):
    def setUp(self):
        self.s, self.sheet = make()
        self.table = tables_ops.create(self.s.book, self.s.stack, self.sheet, (0, 0, 3, 2))

    def test_light_and_dark_families(self):
        from fm_tableur import render
        from fm_tableur.structured import STYLES, style_from_excel
        families = {st["family"] for st in STYLES.values()}
        self.assertEqual(families, {"Clair", "Moyen", "Foncé"})
        tables_ops.set_props(self.s.book, self.s.stack, self.sheet, self.table, style="fonce_vert")
        self.assertEqual(render.table_text_style(self.sheet, 0, 0), {"bold": True, "color": "#ffffff"})
        self.assertEqual(render.table_text_style(self.sheet, 2, 1), {"color": "#ffffff"})   # données en blanc
        tables_ops.set_props(self.s.book, self.s.stack, self.sheet, self.table, style="clair_bleu")
        self.assertEqual(render.table_text_style(self.sheet, 0, 0), {"bold": True})
        self.assertIsNone(render.table_text_style(self.sheet, 2, 1))
        self.assertEqual((style_from_excel("TableStyleDark2"), style_from_excel("TableStyleLight19"),
                          style_from_excel("TableStyleMedium15")), ("fonce_bleu", "clair_bleu", "bleu"))

    def test_unknown_excel_style_kept_until_changed(self):
        import openpyxl
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from fm_tableur import io_xlsx
        tmp = tempfile.mkdtemp()
        try:
            src = os.path.join(tmp, "style.xlsx")
            wb = openpyxl.Workbook()
            for row in [["A", "B"], [1, 2]]:
                wb.active.append(row)
            table = Table(displayName="T", ref="A1:B2")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium15", showRowStripes=True)
            wb.active.add_table(table)
            wb.save(src)
            book = io_xlsx.read(src)
            session = Session(book)
            out = os.path.join(tmp, "out.xlsx")
            io_xlsx.write(book, out)
            self.assertEqual(openpyxl.load_workbook(out).active.tables["T"].tableStyleInfo.name,
                             "TableStyleMedium15")                         # conservé tel quel
            t = book.sheets[0].tables[0]
            tables_ops.set_props(book, session.stack, book.sheets[0], t, style="fonce_or", excel_style=None)
            io_xlsx.write(book, out)
            self.assertEqual(openpyxl.load_workbook(out).active.tables["T"].tableStyleInfo.name, "TableStyleDark5")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_resize_to(self):
        s, sh = self.s, self.sheet
        self.assertIsNone(tables_ops.resize_to(s.book, s.stack, sh, self.table, (0, 0, 5, 3)))
        table = T.find(s.book, "Tableau1")[1]
        self.assertEqual(table["ref"], (0, 0, 5, 3))
        self.assertEqual(sh.cell(0, 3).value, "Colonne4")                 # colonne ajoutée nommée
        put(s, sh, 8, 0, "=ROWS(Tableau1)")
        if fz is not None:
            self.assertEqual(val(s, sh, 8, 0), 5)
        tables_ops.toggle_totals(s.book, s.stack, sh, table)              # totaux en ligne 6
        table = T.find(s.book, "Tableau1")[1]
        self.assertIsNone(tables_ops.resize_to(s.book, s.stack, sh, table, (0, 0, 4, 3)))
        self.assertEqual(sh.cell(4, 0).value, "Total")                    # la ligne des totaux a suivi le bas
        self.assertIsNone(sh.cell(6, 0))
        tables_ops.create(s.book, s.stack, sh, (10, 0, 12, 1))
        self.assertIsNotNone(tables_ops.resize_to(s.book, s.stack, sh, T.find(s.book, "Tableau1")[1],
                                                  (0, 0, 11, 3)))         # chevauchement refusé
        s.stack.undo()                                                    # (création du 2e tableau)
        s.stack.undo()
        self.assertEqual(T.find(s.book, "Tableau1")[1]["ref"], (0, 0, 6, 3))

    def test_handle_drag_in_grid(self):
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtCore import Qt
        from PySide6.QtGui import QMouseEvent
        from PySide6.QtWidgets import QApplication
        app = QApplication.instance() or QApplication([])
        from fm_tableur.widget import SpreadsheetWidget
        w = SpreadsheetWidget(self.s)
        w.resize(900, 500)
        w.show()
        app.processEvents()
        try:
            g = w.grid
            g.select(1, 0)
            corner = g.range_rect(self.table["ref"]).bottomRight()
            start = corner - type(corner)(3, 3)
            self.assertIs(g.table_handle_at(start), self.table)
            vp = g.viewport()
            target = g.cell_rect(6, 3).center()
            for kind, pos, buttons in ((QMouseEvent.MouseButtonPress, start, Qt.LeftButton),
                                       (QMouseEvent.MouseMove, target, Qt.LeftButton),
                                       (QMouseEvent.MouseButtonRelease, target, Qt.NoButton)):
                QApplication.sendEvent(vp, QMouseEvent(kind, pos, vp.mapToGlobal(pos.toPoint()), Qt.LeftButton,
                                                       buttons, Qt.NoModifier))
            self.assertEqual(T.find(self.s.book, "Tableau1")[1]["ref"], (0, 0, 6, 3))
            self.assertEqual(g.view.current, (1, 0))                      # la sélection n'a pas bougé
        finally:
            w.detach()
            w.deleteLater()
